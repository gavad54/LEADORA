"""LEADORA distribution ERP: tenant-scoped sales, inventory, double-entry ledger, treasury and XLSX onboarding."""
import io, json, secrets, sqlite3
from datetime import date
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, Field

LIMIT_BYTES=10*1024*1024

def jalali_to_gregorian(jy,jm,jd):
 jy=int(jy);jm=int(jm);jd=int(jd)
 if jy<1200 or jy>1600 or jm<1 or jm>12 or jd<1 or jd>(31 if jm<=6 else 30 if jm<=11 else 30):
  raise HTTPException(422,"سال یا تاریخ شمسی نامعتبر است.")
 jy+=1595
 days=-355668+365*jy+(jy//33)*8+((jy%33+3)//4)+jd
 if jm<7: days+=(jm-1)*31
 else: days+=(jm-7)*30+186
 gy=400*(days//146097);days%=146097
 if days>36524:
  gy+=100*((days-1)//36524);days=(days-1)%36524
  if days>=365: days+=1
 gy+=4*(days//1461);days%=1461
 if days>365:
  gy+=(days-1)//365;days=(days-1)%365
 gd=days+1;leap=(gy%4==0 and gy%100!=0) or gy%400==0
 md=[31,29 if leap else 28,31,30,31,30,31,31,30,31,30,31];gm=1
 while gm<=12 and gd>md[gm-1]:gd-=md[gm-1];gm+=1
 return date(gy,gm,gd).isoformat()

def jalali_year_range(jy):
 return jalali_to_gregorian(jy,1,1),jalali_to_gregorian(jy+1,1,1)

SCHEMAS={
 "customers":{"title":"مشتریان","required":["کد مشتری","نام مشتری"],"columns":["کد مشتری","نام مشتری","موبایل","تلفن","آدرس","مانده طلب افتتاحیه","توضیحات"]},
 "suppliers":{"title":"تأمین‌کنندگان","required":["کد تأمین‌کننده","نام تأمین‌کننده"],"columns":["کد تأمین‌کننده","نام تأمین‌کننده","موبایل","تلفن","آدرس","مانده بدهی افتتاحیه","توضیحات"]},
 "items":{"title":"کالاها","required":["کد کالا","نام کالا"],"columns":["کد کالا","نام کالا","گروه","واحد","قیمت خرید","قیمت فروش","موجودی اولیه","انبار","محل","حداقل موجودی"]},
 "warehouses":{"title":"انبارها","required":["کد انبار","نام انبار"],"columns":["کد انبار","نام انبار","کد محل","نام محل","محل بالادست"]},
 "banks":{"title":"حساب‌های بانکی","required":["کد بانک","نام بانک"],"columns":["کد بانک","نام بانک","شماره حساب","شماره کارت","شبا","مانده افتتاحیه","توضیحات"]},
 "cash":{"title":"صندوق و تنخواه","required":["کد صندوق","نام صندوق"],"columns":["کد صندوق","نام صندوق","مانده افتتاحیه","مسئول","توضیحات"]},
 "receivables":{"title":"مطالبات افتتاحیه","required":["کد طرف حساب","نام مشتری","مبلغ طلب"],"columns":["کد طرف حساب","نام مشتری","مبلغ طلب","تاریخ سررسید","شرح"]},
 "payables":{"title":"بدهی‌های افتتاحیه","required":["کد طرف حساب","نام تأمین‌کننده","مبلغ بدهی"],"columns":["کد طرف حساب","نام تأمین‌کننده","مبلغ بدهی","تاریخ سررسید","شرح"]},
 "opening_finance":{"title":"مانده افتتاحیه مالی","required":["کد حساب","نام حساب","بدهکار افتتاحیه","بستانکار افتتاحیه"],"columns":["کد حساب","نام حساب","بدهکار افتتاحیه","بستانکار افتتاحیه","شرح"]}}
DEFAULTS=[("1000","صندوق و بانک","asset"),("1010","وجه نقد نزد توزیع‌کننده","asset"),("1020","چک‌های دریافت‌شده در انتظار وصول","asset"),("1100","حساب‌های دریافتنی","asset"),("1200","موجودی کالا","asset"),("2000","حساب‌های پرداختنی","liability"),("3000","سرمایه","equity"),("4000","فروش کالا","income"),("5000","بهای تمام‌شده کالای فروش‌رفته","expense"),("6000","هزینه‌های عملیاتی","expense")]

def minor(v, label="مبلغ"):
 try: d=Decimal(str(v if v not in (None,"") else "0")).quantize(Decimal(".01"),rounding=ROUND_HALF_UP)
 except (InvalidOperation,ValueError,TypeError): raise HTTPException(422,f"{label} نامعتبر است.")
 if not d.is_finite() or abs(d)>Decimal("1000000000000"): raise HTTPException(422,f"{label} خارج از محدوده مجاز است.")
 return int(d*100)
def money(v): return format(Decimal(int(v or 0))/100,".2f")
def audit(c,t,action): c.execute("INSERT INTO audit_log(tenant_id,action) VALUES(?,?)",(t,action))
def init_db(db):
 with db() as c:
  c.executescript("""
  CREATE TABLE IF NOT EXISTS dist_accounts(id INTEGER PRIMARY KEY,tenant_id INTEGER NOT NULL,code TEXT NOT NULL,name TEXT NOT NULL,nature TEXT NOT NULL,parent_code TEXT NOT NULL DEFAULT '',active INTEGER NOT NULL DEFAULT 1,UNIQUE(tenant_id,code));
  CREATE TABLE IF NOT EXISTS dist_journals(id INTEGER PRIMARY KEY,tenant_id INTEGER NOT NULL,entry_no TEXT NOT NULL,entry_date TEXT NOT NULL,description TEXT NOT NULL,source_type TEXT NOT NULL,source_id INTEGER,status TEXT NOT NULL DEFAULT 'POSTED',created_by INTEGER,created_at TEXT DEFAULT CURRENT_TIMESTAMP,UNIQUE(tenant_id,entry_no));
  CREATE TABLE IF NOT EXISTS dist_journal_lines(id INTEGER PRIMARY KEY,tenant_id INTEGER NOT NULL,journal_id INTEGER NOT NULL REFERENCES dist_journals(id),account_code TEXT NOT NULL,debit_minor INTEGER NOT NULL DEFAULT 0 CHECK(debit_minor>=0),credit_minor INTEGER NOT NULL DEFAULT 0 CHECK(credit_minor>=0),party_code TEXT NOT NULL DEFAULT '',memo TEXT NOT NULL DEFAULT '',CHECK((debit_minor=0)!=(credit_minor=0)));
  CREATE TABLE IF NOT EXISTS dist_customer_codes(tenant_id INTEGER NOT NULL,customer_id INTEGER NOT NULL,code TEXT NOT NULL,UNIQUE(tenant_id,code),UNIQUE(tenant_id,customer_id));
  CREATE TABLE IF NOT EXISTS dist_suppliers(id INTEGER PRIMARY KEY,tenant_id INTEGER NOT NULL,code TEXT NOT NULL,name TEXT NOT NULL,phone TEXT NOT NULL DEFAULT '',address TEXT NOT NULL DEFAULT '',UNIQUE(tenant_id,code));
  CREATE TABLE IF NOT EXISTS dist_warehouses(id INTEGER PRIMARY KEY,tenant_id INTEGER NOT NULL,code TEXT NOT NULL,name TEXT NOT NULL,location TEXT NOT NULL DEFAULT '',UNIQUE(tenant_id,code));
  CREATE TABLE IF NOT EXISTS dist_item_data(tenant_id INTEGER NOT NULL,product_id INTEGER NOT NULL,cost_minor INTEGER NOT NULL DEFAULT 0,min_stock REAL NOT NULL DEFAULT 0,unit TEXT NOT NULL DEFAULT 'عدد',warehouse_code TEXT NOT NULL DEFAULT 'DEFAULT',PRIMARY KEY(tenant_id,product_id));
  CREATE TABLE IF NOT EXISTS dist_stock_moves(id INTEGER PRIMARY KEY,tenant_id INTEGER NOT NULL,product_id INTEGER NOT NULL,warehouse_code TEXT NOT NULL,quantity REAL NOT NULL,cost_minor INTEGER NOT NULL DEFAULT 0,kind TEXT NOT NULL,reference TEXT NOT NULL DEFAULT '',created_at TEXT DEFAULT CURRENT_TIMESTAMP,created_by INTEGER);
  CREATE TABLE IF NOT EXISTS dist_sales(id INTEGER PRIMARY KEY,tenant_id INTEGER NOT NULL,order_no TEXT NOT NULL,customer_id INTEGER NOT NULL,order_date TEXT NOT NULL,status TEXT NOT NULL DEFAULT 'DRAFT',total_minor INTEGER NOT NULL DEFAULT 0,note TEXT NOT NULL DEFAULT '',UNIQUE(tenant_id,order_no));
  CREATE TABLE IF NOT EXISTS dist_sale_lines(id INTEGER PRIMARY KEY,tenant_id INTEGER NOT NULL,order_id INTEGER NOT NULL,product_id INTEGER NOT NULL,quantity REAL NOT NULL,unit_price_minor INTEGER NOT NULL,discount_minor INTEGER NOT NULL DEFAULT 0,total_minor INTEGER NOT NULL);
  CREATE TABLE IF NOT EXISTS dist_purchases(id INTEGER PRIMARY KEY,tenant_id INTEGER NOT NULL,order_no TEXT NOT NULL,supplier_id INTEGER NOT NULL,warehouse_code TEXT NOT NULL,order_date TEXT NOT NULL,status TEXT NOT NULL DEFAULT 'DRAFT',total_minor INTEGER NOT NULL DEFAULT 0,UNIQUE(tenant_id,order_no));
  CREATE TABLE IF NOT EXISTS dist_purchase_lines(id INTEGER PRIMARY KEY,tenant_id INTEGER NOT NULL,order_id INTEGER NOT NULL,product_id INTEGER NOT NULL,quantity REAL NOT NULL,cost_minor INTEGER NOT NULL,total_minor INTEGER NOT NULL);
  CREATE TABLE IF NOT EXISTS dist_deliveries(id INTEGER PRIMARY KEY,tenant_id INTEGER NOT NULL,order_id INTEGER NOT NULL,reference TEXT NOT NULL DEFAULT '',note TEXT NOT NULL DEFAULT '',created_at TEXT DEFAULT CURRENT_TIMESTAMP,UNIQUE(tenant_id,order_id));
  CREATE TABLE IF NOT EXISTS dist_cash_accounts(id INTEGER PRIMARY KEY,tenant_id INTEGER NOT NULL,code TEXT NOT NULL,name TEXT NOT NULL,kind TEXT NOT NULL,opening_minor INTEGER NOT NULL DEFAULT 0,UNIQUE(tenant_id,code));
  CREATE TABLE IF NOT EXISTS dist_cash_moves(id INTEGER PRIMARY KEY,tenant_id INTEGER NOT NULL,account_id INTEGER NOT NULL,txn_date TEXT NOT NULL,direction TEXT NOT NULL,amount_minor INTEGER NOT NULL,reference TEXT NOT NULL DEFAULT '',memo TEXT NOT NULL DEFAULT '');
  CREATE TABLE IF NOT EXISTS dist_customer_receipts(id INTEGER PRIMARY KEY,tenant_id INTEGER NOT NULL,order_id INTEGER NOT NULL,account_id INTEGER NOT NULL,txn_date TEXT NOT NULL,amount_minor INTEGER NOT NULL,reference TEXT NOT NULL DEFAULT '');
  CREATE TABLE IF NOT EXISTS dist_supplier_payments(id INTEGER PRIMARY KEY,tenant_id INTEGER NOT NULL,purchase_id INTEGER NOT NULL,account_id INTEGER NOT NULL,txn_date TEXT NOT NULL,amount_minor INTEGER NOT NULL,reference TEXT NOT NULL DEFAULT '');
  CREATE TABLE IF NOT EXISTS dist_import_batches(id TEXT PRIMARY KEY,tenant_id INTEGER NOT NULL,dataset TEXT NOT NULL,status TEXT NOT NULL DEFAULT 'PREVIEW',rows_json TEXT NOT NULL,report_json TEXT NOT NULL,created_at TEXT DEFAULT CURRENT_TIMESTAMP,committed_at TEXT);
  CREATE TABLE IF NOT EXISTS dist_opening_balances(id INTEGER PRIMARY KEY,tenant_id INTEGER NOT NULL,kind TEXT NOT NULL,code TEXT NOT NULL,name TEXT NOT NULL,amount_minor INTEGER NOT NULL,due_date TEXT NOT NULL DEFAULT '',description TEXT NOT NULL DEFAULT '',batch_id TEXT NOT NULL);
  CREATE TABLE IF NOT EXISTS dist_period_locks(tenant_id INTEGER NOT NULL,period TEXT NOT NULL,PRIMARY KEY(tenant_id,period));
  CREATE TABLE IF NOT EXISTS dist_fiscal_years(id INTEGER PRIMARY KEY,tenant_id INTEGER NOT NULL,jalali_year INTEGER NOT NULL,start_date TEXT NOT NULL,end_date TEXT NOT NULL,status TEXT NOT NULL DEFAULT 'OPEN',carry_forward_done INTEGER NOT NULL DEFAULT 0,created_at TEXT DEFAULT CURRENT_TIMESTAMP,closed_at TEXT,UNIQUE(tenant_id,jalali_year));
  CREATE TABLE IF NOT EXISTS dist_fiscal_carryovers(id INTEGER PRIMARY KEY,tenant_id INTEGER NOT NULL,from_year_id INTEGER NOT NULL,to_year_id INTEGER NOT NULL,account_code TEXT NOT NULL,amount_minor INTEGER NOT NULL,created_at TEXT DEFAULT CURRENT_TIMESTAMP,UNIQUE(tenant_id,from_year_id,to_year_id,account_code));
  """)
  c.executescript("""
  CREATE TABLE IF NOT EXISTS dist_sales_returns(
    id INTEGER PRIMARY KEY, tenant_id INTEGER NOT NULL, order_id INTEGER NOT NULL,
    return_no TEXT NOT NULL, return_date TEXT NOT NULL, reason TEXT NOT NULL DEFAULT '',
    status TEXT NOT NULL DEFAULT 'DRAFT' CHECK(status IN ('DRAFT','APPROVED','REJECTED')),
    created_by INTEGER, created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP, approved_at TEXT,
    UNIQUE(tenant_id,return_no)
  );
  CREATE TABLE IF NOT EXISTS dist_sales_return_lines(
    id INTEGER PRIMARY KEY, tenant_id INTEGER NOT NULL,
    return_id INTEGER NOT NULL REFERENCES dist_sales_returns(id),
    product_id INTEGER NOT NULL, quantity REAL NOT NULL CHECK(quantity>0),
    revenue_minor INTEGER NOT NULL DEFAULT 0 CHECK(revenue_minor>=0),
    cost_minor INTEGER NOT NULL DEFAULT 0 CHECK(cost_minor>=0),
    UNIQUE(tenant_id,return_id,product_id)
  );
  CREATE TABLE IF NOT EXISTS dist_dispatch_manifests(
    id INTEGER PRIMARY KEY,
    tenant_id INTEGER NOT NULL,
    manifest_no TEXT NOT NULL,
    dispatch_date TEXT NOT NULL,
    distributor_name TEXT NOT NULL,
    status TEXT NOT NULL DEFAULT 'OPEN' CHECK(status IN ('OPEN','SETTLED','CANCELLED')),
    created_by INTEGER,
    created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
    UNIQUE(tenant_id,manifest_no)
  );
  CREATE TABLE IF NOT EXISTS dist_dispatch_manifest_lines(
    id INTEGER PRIMARY KEY,
    tenant_id INTEGER NOT NULL,
    manifest_id INTEGER NOT NULL REFERENCES dist_dispatch_manifests(id),
    order_id INTEGER NOT NULL,
    invoice_total_minor INTEGER NOT NULL CHECK(invoice_total_minor>=0),
    cash_minor INTEGER NOT NULL DEFAULT 0 CHECK(cash_minor>=0),
    check_minor INTEGER NOT NULL DEFAULT 0 CHECK(check_minor>=0),
    credit_minor INTEGER NOT NULL DEFAULT 0 CHECK(credit_minor>=0),
    UNIQUE(tenant_id,order_id)
  );
  CREATE INDEX IF NOT EXISTS idx_dispatch_manifests_tenant_date
    ON dist_dispatch_manifests(tenant_id,dispatch_date,id);
  CREATE TABLE IF NOT EXISTS dist_dispatch_transactions(
    id INTEGER PRIMARY KEY, tenant_id INTEGER NOT NULL, manifest_id INTEGER NOT NULL,
    order_id INTEGER, transaction_type TEXT NOT NULL CHECK(transaction_type IN ('CASH_COLLECTION','CHECK_COLLECTION','CASH_REMITTANCE')),
    amount_minor INTEGER NOT NULL CHECK(amount_minor>0), txn_date TEXT NOT NULL,
    reference TEXT NOT NULL DEFAULT '', account_id INTEGER, status TEXT NOT NULL DEFAULT 'RECORDED',
    created_by INTEGER, created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
  );
  CREATE INDEX IF NOT EXISTS idx_dispatch_transactions_tenant_manifest
    ON dist_dispatch_transactions(tenant_id,manifest_id,transaction_type,status);
  """)
  cols={x["name"] for x in c.execute("PRAGMA table_info(dist_purchases)").fetchall()}
  if "warehouse_code" not in cols: c.execute("ALTER TABLE dist_purchases ADD COLUMN warehouse_code TEXT NOT NULL DEFAULT 'DEFAULT'")

def post(c,t,dt,desc,source,source_id,lines,uid=None):
 try: date.fromisoformat(dt)
 except (ValueError,TypeError): raise HTTPException(422,"تاریخ باید YYYY-MM-DD باشد.")
 if c.execute("SELECT 1 FROM dist_period_locks WHERE tenant_id=? AND period=?",(t,dt[:7])).fetchone(): raise HTTPException(409,"دوره مالی بسته است.")
 fy=c.execute("SELECT status FROM dist_fiscal_years WHERE tenant_id=? AND start_date<=? AND end_date>=?",(t,dt,dt)).fetchone()
 if fy and fy["status"]=="CLOSED":raise HTTPException(409,"سال مالی بسته شده است؛ ثبت سند در این سال مجاز نیست.")
 clean=[]
 for x in lines:
  code=str(x["code"]); d=int(x.get("d",0)); cr=int(x.get("c",0))
  if not c.execute("SELECT 1 FROM dist_accounts WHERE tenant_id=? AND code=? AND active=1",(t,code)).fetchone(): raise HTTPException(422,f"حساب {code} تعریف نشده است.")
  if d<0 or cr<0 or (d==0)==(cr==0): raise HTTPException(422,"هر ردیف سند باید فقط بدهکار یا بستانکار باشد.")
  clean.append((code,d,cr))
 if len(clean)<2 or sum(x[1] for x in clean)!=sum(x[2] for x in clean): raise HTTPException(422,"سند نامتوازن است.")
 no=f"{dt[:4]}-{secrets.token_hex(4).upper()}"
 jid=c.execute("INSERT INTO dist_journals(tenant_id,entry_no,entry_date,description,source_type,source_id,created_by) VALUES(?,?,?,?,?,?,?)",(t,no,dt,desc,source,source_id,uid)).lastrowid
 c.executemany("INSERT INTO dist_journal_lines(tenant_id,journal_id,account_code,debit_minor,credit_minor) VALUES(?,?,?,?,?)",[(t,jid,*x) for x in clean])
 return no

def dispatch_sale(c,t,oid,dispatch_date,reference,note,uid=None):
 try:date.fromisoformat(dispatch_date)
 except (ValueError,TypeError):raise HTTPException(422,"تاریخ ارسال باید YYYY-MM-DD باشد.")
 o=c.execute("SELECT order_no,order_date,status FROM dist_sales WHERE tenant_id=? AND id=?",(t,oid)).fetchone()
 if not o:raise HTTPException(404,"فاکتور در شرکت شما یافت نشد.")
 if o["status"]!="POSTED":raise HTTPException(409,"فقط فاکتور نهایی‌شده و ارسال‌نشده قابل حواله است.")
 if c.execute("SELECT 1 FROM dist_deliveries WHERE tenant_id=? AND order_id=?",(t,oid)).fetchone():raise HTTPException(409,"برای این فاکتور قبلاً حواله ثبت شده است.")
 lines=c.execute("SELECT l.*,p.stock,COALESCE(d.cost_minor,0) cost,COALESCE(d.warehouse_code,'DEFAULT') wh FROM dist_sale_lines l JOIN products p ON p.id=l.product_id AND p.tenant_id=l.tenant_id LEFT JOIN dist_item_data d ON d.tenant_id=p.tenant_id AND d.product_id=p.id WHERE l.tenant_id=? AND l.order_id=?",(t,oid)).fetchall()
 for x in lines:
  if x["stock"]<x["quantity"]:raise HTTPException(409,f"موجودی کالای {x['product_id']} برای ارسال کافی نیست.")
 cost=sum(minor(Decimal(str(x["quantity"]))*Decimal(x["cost"])/100) for x in lines)
 if cost:post(c,t,dispatch_date,"بهای تمام‌شده ارسال "+o["order_no"],"COGS",oid,[{"code":"5000","d":cost},{"code":"1200","c":cost}],uid)
 for x in lines:
  changed=c.execute("UPDATE products SET stock=stock-? WHERE tenant_id=? AND id=? AND stock>=?",(x["quantity"],t,x["product_id"],x["quantity"]))
  if changed.rowcount!=1:raise HTTPException(409,f"موجودی کالای {x['product_id']} هم‌زمان تغییر کرده است؛ دوباره تلاش کنید.")
  c.execute("INSERT INTO dist_stock_moves(tenant_id,product_id,warehouse_code,quantity,cost_minor,kind,reference,created_by) VALUES(?,?,?,?,?,?,?,?)",(t,x["product_id"],x["wh"],x["quantity"],x["cost"],"SALE_OUT",o["order_no"],uid))
 c.execute("INSERT INTO dist_deliveries(tenant_id,order_id,reference,note) VALUES(?,?,?,?)",(t,oid,reference,note))
 c.execute("UPDATE dist_sales SET status='DELIVERED' WHERE tenant_id=? AND id=? AND status='POSTED'",(t,oid))
 audit(c,t,"distribution.sale.dispatched:"+o["order_no"])
 return {"status":"DELIVERED","order_no":o["order_no"],"dispatch_date":dispatch_date,"cost_of_goods":money(cost),"inventory_deducted":True}

class Warehouse(BaseModel):
 code:str=Field(min_length=1,max_length=40); name:str=Field(min_length=1,max_length=160); location:str=""

class FiscalYearIn(BaseModel):
 jalali_year:int=Field(ge=1200,le=1599)

class SaleLine(BaseModel):
 product_id:int; quantity:float=Field(gt=0,le=1e9); unit_price:Decimal=Field(ge=0); discount:Decimal=Field(default=Decimal("0"),ge=0)
class Sale(BaseModel):
 order_no:str=Field(min_length=1,max_length=80); customer_id:int; order_date:str=Field(min_length=10,max_length=10); lines:list[SaleLine]=Field(min_length=1,max_length=200); note:str=""
class PurchaseLine(BaseModel):
 product_id:int; quantity:float=Field(gt=0,le=1e9); unit_cost:Decimal=Field(ge=0)
class Purchase(BaseModel):
 order_no:str; supplier_id:int; warehouse_code:str; order_date:str=Field(min_length=10,max_length=10); lines:list[PurchaseLine]=Field(min_length=1,max_length=200)
class CashAccount(BaseModel):
 code:str; name:str; kind:str=Field(pattern="^(CASH|BANK)$"); opening_balance:Decimal=Decimal("0")
class CashMove(BaseModel):
 code:str; txn_date:str=Field(min_length=10,max_length=10); direction:str=Field(pattern="^(IN|OUT)$"); amount:Decimal=Field(gt=0); reference:str=""; memo:str=""
class ReceiptIn(BaseModel):
 order_id:int; cash_code:str; txn_date:str=Field(min_length=10,max_length=10); amount:Decimal=Field(gt=0); reference:str=""
class DispatchAllocation(BaseModel):
 order_id:int=Field(gt=0)
 cash_amount:Decimal=Field(default=Decimal("0"),ge=0)
 check_amount:Decimal=Field(default=Decimal("0"),ge=0)
 credit_amount:Decimal=Field(default=Decimal("0"),ge=0)

class DispatchManifestIn(BaseModel):
 manifest_no:str=Field(min_length=1,max_length=80)
 dispatch_date:str=Field(min_length=10,max_length=10)
 distributor_name:str=Field(min_length=1,max_length=160)
 invoices:list[DispatchAllocation]=Field(min_length=1,max_length=500)

class DispatchCollectionIn(BaseModel):
 order_id:int=Field(gt=0)
 payment_type:str=Field(pattern="^(CASH|CHECK)$")
 amount:Decimal=Field(gt=0)
 txn_date:str=Field(min_length=10,max_length=10)
 reference:str=Field(default="",max_length=160)

class DispatchRemittanceIn(BaseModel):
 cash_code:str=Field(min_length=1,max_length=40)
 amount:Decimal=Field(gt=0)
 txn_date:str=Field(min_length=10,max_length=10)
 reference:str=Field(default="",max_length=160)

class SalesReturnLineIn(BaseModel):
 product_id:int=Field(gt=0)
 quantity:Decimal=Field(gt=0,le=Decimal("1000000000"))

class SalesReturnIn(BaseModel):
 order_id:int=Field(gt=0)
 return_no:str=Field(min_length=1,max_length=80)
 return_date:str=Field(min_length=10,max_length=10)
 reason:str=Field(default="",max_length=500)
 lines:list[SalesReturnLineIn]=Field(min_length=1,max_length=200)

class AccountIn(BaseModel):
 code:str=Field(min_length=1,max_length=30); name:str=Field(min_length=1,max_length=160); nature:str=Field(pattern="^(asset|liability|equity|income|expense)$"); parent_code:str=""
class JournalLineIn(BaseModel):
 account_code:str; debit:Decimal=Decimal("0"); credit:Decimal=Decimal("0")
class JournalIn(BaseModel):
 entry_date:str=Field(min_length=10,max_length=10); description:str=Field(min_length=1,max_length=240); lines:list[JournalLineIn]=Field(min_length=2,max_length=200)

def create_router(db,user_dep):
 init_db(db); r=APIRouter(prefix="/api/distribution",tags=["distribution-erp"])
 def tenant(u):
  t=int(u["tenant_id"])
  with db() as c:
   for code,name,nature in DEFAULTS:c.execute("INSERT OR IGNORE INTO dist_accounts(tenant_id,code,name,nature) VALUES(?,?,?,?)",(t,code,name,nature))
  return t
 @r.get("/fiscal-years")
 def fiscal_years(u=Depends(user_dep)):
  with db() as c:
   rows=c.execute("SELECT id,jalali_year,start_date,end_date,status,carry_forward_done,created_at,closed_at FROM dist_fiscal_years WHERE tenant_id=? ORDER BY jalali_year DESC",(tenant(u),)).fetchall()
  return [dict(x) for x in rows]
 @r.post("/fiscal-years")
 def create_fiscal_year(b:FiscalYearIn,u=Depends(user_dep)):
  t=tenant(u);start,next_start=jalali_year_range(b.jalali_year);end=(date.fromisoformat(next_start)-__import__("datetime").timedelta(days=1)).isoformat()
  with db() as c:
   try:c.execute("INSERT INTO dist_fiscal_years(tenant_id,jalali_year,start_date,end_date) VALUES(?,?,?,?)",(t,b.jalali_year,start,end))
   except sqlite3.IntegrityError:raise HTTPException(409,"این سال مالی قبلاً تعریف شده است.")
   row=c.execute("SELECT id,jalali_year,start_date,end_date,status,carry_forward_done FROM dist_fiscal_years WHERE tenant_id=? AND jalali_year=?",(t,b.jalali_year)).fetchone();audit(c,t,"distribution.fiscal_year.created:"+str(b.jalali_year))
  return dict(row)
 @r.post("/fiscal-years/{fyid}/close")
 def close_fiscal_year(fyid:int,u=Depends(user_dep)):
  t=tenant(u)
  with db() as c:
   fy=c.execute("SELECT * FROM dist_fiscal_years WHERE tenant_id=? AND id=?",(t,fyid)).fetchone()
   if not fy:raise HTTPException(404,"سال مالی یافت نشد.")
   if fy["status"]=="CLOSED":return {"status":"CLOSED","jalali_year":fy["jalali_year"]}
   unbalanced=c.execute("SELECT j.entry_no FROM dist_journals j JOIN dist_journal_lines l ON l.tenant_id=j.tenant_id AND l.journal_id=j.id WHERE j.tenant_id=? AND j.entry_date BETWEEN ? AND ? AND j.status='POSTED' GROUP BY j.id HAVING SUM(l.debit_minor)!=SUM(l.credit_minor) LIMIT 1",(t,fy["start_date"],fy["end_date"])).fetchone()
   if unbalanced:raise HTTPException(409,"سند نامتوازن وجود دارد: "+unbalanced["entry_no"])
   c.execute("UPDATE dist_fiscal_years SET status='CLOSED',closed_at=CURRENT_TIMESTAMP WHERE tenant_id=? AND id=?",(t,fyid))
   audit(c,t,"distribution.fiscal_year.closed:"+str(fy["jalali_year"]))
  return {"status":"CLOSED","jalali_year":fy["jalali_year"]}
 @r.post("/fiscal-years/{fyid}/carry-forward")
 def carry_forward(fyid:int,u=Depends(user_dep)):
  t=tenant(u)
  with db() as c:
   fy=c.execute("SELECT * FROM dist_fiscal_years WHERE tenant_id=? AND id=?",(t,fyid)).fetchone()
   if not fy:raise HTTPException(404,"سال مالی یافت نشد.")
   if fy["status"]!="CLOSED":raise HTTPException(409,"ابتدا سال مالی را ببندید.")
   nxt=c.execute("SELECT * FROM dist_fiscal_years WHERE tenant_id=? AND jalali_year=?",(t,fy["jalali_year"]+1)).fetchone()
   if not nxt:raise HTTPException(409,"ابتدا سال مالی بعد را تعریف کنید.")
   if fy["carry_forward_done"]:raise HTTPException(409,"انتقال مانده‌ها قبلاً انجام شده است.")
   rows=c.execute("SELECT a.code,COALESCE(SUM(l.debit_minor-l.credit_minor),0) bal FROM dist_accounts a LEFT JOIN dist_journals j ON j.tenant_id=a.tenant_id AND j.entry_date BETWEEN ? AND ? AND j.status='POSTED' LEFT JOIN dist_journal_lines l ON l.tenant_id=j.tenant_id AND l.journal_id=j.id AND l.account_code=a.code WHERE a.tenant_id=? AND a.nature IN ('asset','liability','equity') GROUP BY a.code",(fy["start_date"],fy["end_date"],t)).fetchall()
   for x in rows:c.execute("INSERT INTO dist_fiscal_carryovers(tenant_id,from_year_id,to_year_id,account_code,amount_minor) VALUES(?,?,?,?,?)",(t,fyid,nxt["id"],x["code"],x["bal"]))
   c.execute("UPDATE dist_fiscal_years SET carry_forward_done=1 WHERE tenant_id=? AND id=?",(t,fyid));audit(c,t,"distribution.fiscal_year.carry_forward:"+str(fy["jalali_year"]))
  return {"status":"TRANSFERRED","from_year":fy["jalali_year"],"to_year":fy["jalali_year"]+1,"accounts_transferred":len(rows),"note":"سوابق اسناد حفظ شده‌اند؛ انتقال مانده ترازنامه به‌صورت snapshot قابل حسابرسی ثبت شد."}
 @r.get("/fiscal-years/{fyid}/summary")
 def fiscal_summary(fyid:int,u=Depends(user_dep)):
  t=tenant(u)
  with db() as c:
   fy=c.execute("SELECT * FROM dist_fiscal_years WHERE tenant_id=? AND id=?",(t,fyid)).fetchone()
   if not fy:raise HTTPException(404,"سال مالی یافت نشد.")
   sales=c.execute("SELECT COUNT(*) n,COALESCE(SUM(total_minor),0) total FROM dist_sales WHERE tenant_id=? AND order_date BETWEEN ? AND ? AND status!='CANCELLED'",(t,fy["start_date"],fy["end_date"])).fetchone()
   journals=c.execute("SELECT COUNT(*) FROM dist_journals WHERE tenant_id=? AND entry_date BETWEEN ? AND ? AND status='POSTED'",(t,fy["start_date"],fy["end_date"])).fetchone()[0]
  return {"jalali_year":fy["jalali_year"],"start_date":fy["start_date"],"end_date":fy["end_date"],"status":fy["status"],"sales_orders":sales["n"],"sales_total":money(sales["total"]),"posted_journals":journals,"carry_forward_done":bool(fy["carry_forward_done"])}
 @r.get("/overview")
 def overview(u=Depends(user_dep)):
  t=tenant(u)
  with db() as c:
   p=c.execute("SELECT COUNT(*) n,COALESCE(SUM(stock),0) stock FROM products WHERE tenant_id=?",(t,)).fetchone()
   cust=c.execute("SELECT COUNT(*) FROM customers WHERE tenant_id=?",(t,)).fetchone()[0]
   s=c.execute("SELECT COUNT(*),COALESCE(SUM(total_minor),0) FROM dist_sales WHERE tenant_id=? AND status!='CANCELLED'",(t,)).fetchone()
   j=c.execute("SELECT COUNT(*) FROM dist_journals WHERE tenant_id=? AND status='POSTED'",(t,)).fetchone()[0]
   low=c.execute("SELECT p.sku,p.name,p.stock,COALESCE(d.min_stock,5) min_stock FROM products p LEFT JOIN dist_item_data d ON d.tenant_id=p.tenant_id AND d.product_id=p.id WHERE p.tenant_id=? AND p.stock<=COALESCE(d.min_stock,5) ORDER BY p.stock LIMIT 20",(t,)).fetchall()
  return {"customers":cust,"products":p["n"],"stock_units":p["stock"],"sales_orders":s[0],"sales_total":money(s[1]),"posted_journals":j,"low_stock":[dict(x) for x in low]}
 @r.get("/accounts")
 def accounts(u=Depends(user_dep)):
  with db() as c:return [dict(x) for x in c.execute("SELECT code,name,nature,parent_code FROM dist_accounts WHERE tenant_id=? AND active=1 ORDER BY code",(tenant(u),)).fetchall()]
 @r.get("/warehouses")
 def warehouses(u=Depends(user_dep)):
  with db() as c:return [dict(x) for x in c.execute("SELECT code,name,location FROM dist_warehouses WHERE tenant_id=? ORDER BY code",(tenant(u),)).fetchall()]
 @r.post("/warehouses")
 def add_warehouse(b:Warehouse,u=Depends(user_dep)):
  t=tenant(u)
  with db() as c:
   try:c.execute("INSERT INTO dist_warehouses(tenant_id,code,name,location) VALUES(?,?,?,?)",(t,b.code.strip(),b.name.strip(),b.location))
   except sqlite3.IntegrityError:raise HTTPException(409,"کد انبار تکراری است.")
   audit(c,t,"distribution.warehouse.created:"+b.code)
  return {"status":"created","code":b.code}
 @r.post("/dispatch/manifests")
 def create_dispatch_manifest(b:DispatchManifestIn,u=Depends(user_dep)):
  t=tenant(u)
  if not b.manifest_no.strip() or not b.distributor_name.strip():raise HTTPException(422,"شماره مانیفست و نام توزیع‌کننده الزامی است.")
  try:date.fromisoformat(b.dispatch_date)
  except (ValueError,TypeError):raise HTTPException(422,"تاریخ حواله باید YYYY-MM-DD باشد.")
  order_ids=[x.order_id for x in b.invoices]
  if len(set(order_ids))!=len(order_ids):raise HTTPException(422,"یک فاکتور در مانیفست تکرار شده است.")
  allocations=[]
  total=cash_total=check_total=credit_total=0
  with db() as c:
   for x in b.invoices:
    order=c.execute("SELECT id,order_no,customer_id,status,total_minor FROM dist_sales WHERE tenant_id=? AND id=?",(t,x.order_id)).fetchone()
    if not order:raise HTTPException(404,"فاکتور در شرکت شما یافت نشد.")
    if order["status"]!="POSTED":raise HTTPException(409,"فقط فاکتور نهایی‌شده و هنوز ارسال‌نشده قابل درج در مانیفست است.")
    received=c.execute("SELECT COALESCE(SUM(amount_minor),0) FROM dist_customer_receipts WHERE tenant_id=? AND order_id=?",(t,x.order_id)).fetchone()[0]
    outstanding=order["total_minor"]-received
    if outstanding<0:raise HTTPException(409,"جمع دریافت‌های ثبت‌شده از مبلغ فاکتور بیشتر است.")
    cash=minor(x.cash_amount,"مبلغ نقد")
    check=minor(x.check_amount,"مبلغ چک")
    credit=minor(x.credit_amount,"مبلغ اعتباری")
    if cash+check+credit!=outstanding:raise HTTPException(422,f"تفکیک نقد/چک/اعتباری فاکتور {order['order_no']} با مانده فاکتور برابر نیست.")
    if c.execute("SELECT 1 FROM dist_dispatch_manifest_lines WHERE tenant_id=? AND order_id=?",(t,x.order_id)).fetchone():
     raise HTTPException(409,"این فاکتور قبلاً در یک مانیفست ثبت شده است.")
    allocations.append((x.order_id,outstanding,cash,check,credit))
    total+=outstanding;cash_total+=cash;check_total+=check;credit_total+=credit
   try:
    mid=c.execute("INSERT INTO dist_dispatch_manifests(tenant_id,manifest_no,dispatch_date,distributor_name,created_by) VALUES(?,?,?,?,?)",
      (t,b.manifest_no.strip(),b.dispatch_date,b.distributor_name.strip(),u.get("id"))).lastrowid
   except sqlite3.IntegrityError:raise HTTPException(409,"شماره مانیفست تکراری است.")
   try:c.executemany("INSERT INTO dist_dispatch_manifest_lines(tenant_id,manifest_id,order_id,invoice_total_minor,cash_minor,check_minor,credit_minor) VALUES(?,?,?,?,?,?,?)",[(t,mid,*x) for x in allocations])
   except sqlite3.IntegrityError:raise HTTPException(409,"یکی از فاکتورها هم‌زمان در مانیفست دیگری ثبت شده است.")
   audit(c,t,"distribution.dispatch_manifest.created:"+b.manifest_no.strip())
  return {"id":mid,"manifest_no":b.manifest_no.strip(),"dispatch_date":b.dispatch_date,"distributor_name":b.distributor_name.strip(),
          "status":"OPEN","invoice_count":len(allocations),"total":money(total),"cash_total":money(cash_total),
          "check_total":money(check_total),"credit_total":money(credit_total)}

 @r.get("/dispatch/manifests")
 def list_dispatch_manifests(u=Depends(user_dep)):
  t=tenant(u)
  with db() as c:
   rows=c.execute("""SELECT m.id,m.manifest_no,m.dispatch_date,m.distributor_name,m.status,m.created_at,
      COUNT(l.id) invoice_count,COALESCE(SUM(l.invoice_total_minor),0) total,
      COALESCE(SUM(l.cash_minor),0) cash_total,COALESCE(SUM(l.check_minor),0) check_total,
      COALESCE(SUM(l.credit_minor),0) credit_total
      FROM dist_dispatch_manifests m LEFT JOIN dist_dispatch_manifest_lines l
      ON l.tenant_id=m.tenant_id AND l.manifest_id=m.id
      WHERE m.tenant_id=? GROUP BY m.id ORDER BY m.id DESC LIMIT 500""",(t,)).fetchall()
  return [{"id":x["id"],"manifest_no":x["manifest_no"],"dispatch_date":x["dispatch_date"],
    "distributor_name":x["distributor_name"],"status":x["status"],"created_at":x["created_at"],
    "invoice_count":x["invoice_count"],"total":money(x["total"]),"cash_total":money(x["cash_total"]),
    "check_total":money(x["check_total"]),"credit_total":money(x["credit_total"])} for x in rows]

 @r.get("/dispatch/manifests/{manifest_id}")
 def get_dispatch_manifest(manifest_id:int,u=Depends(user_dep)):
  t=tenant(u)
  with db() as c:
   m=c.execute("""SELECT m.id,m.manifest_no,m.dispatch_date,m.distributor_name,m.status,m.created_at,
      COUNT(l.id) invoice_count,COALESCE(SUM(l.invoice_total_minor),0) total,
      COALESCE(SUM(l.cash_minor),0) cash_total,COALESCE(SUM(l.check_minor),0) check_total,
      COALESCE(SUM(l.credit_minor),0) credit_total
      FROM dist_dispatch_manifests m LEFT JOIN dist_dispatch_manifest_lines l
      ON l.tenant_id=m.tenant_id AND l.manifest_id=m.id
      WHERE m.tenant_id=? AND m.id=? GROUP BY m.id""",(t,manifest_id)).fetchone()
   if not m:raise HTTPException(404,"مانیفست در شرکت شما یافت نشد.")
   lines=c.execute("""SELECT l.order_id,s.order_no,s.order_date,s.status invoice_status,cu.name customer_name,
      l.invoice_total_minor,l.cash_minor,l.check_minor,l.credit_minor
      FROM dist_dispatch_manifest_lines l
      JOIN dist_sales s ON s.id=l.order_id AND s.tenant_id=l.tenant_id
      JOIN customers cu ON cu.id=s.customer_id AND cu.tenant_id=s.tenant_id
      WHERE l.tenant_id=? AND l.manifest_id=? ORDER BY l.id""",(t,manifest_id)).fetchall()
  return {"id":m["id"],"manifest_no":m["manifest_no"],"dispatch_date":m["dispatch_date"],
    "distributor_name":m["distributor_name"],"status":m["status"],"created_at":m["created_at"],
    "invoice_count":m["invoice_count"],"total":money(m["total"]),"cash_total":money(m["cash_total"]),
    "check_total":money(m["check_total"]),"credit_total":money(m["credit_total"]),
    "invoices":[{"order_id":x["order_id"],"order_no":x["order_no"],"order_date":x["order_date"],
      "invoice_status":x["invoice_status"],"customer_name":x["customer_name"],
      "invoice_total":money(x["invoice_total_minor"]),"cash_amount":money(x["cash_minor"]),
      "check_amount":money(x["check_minor"]),"credit_amount":money(x["credit_minor"])} for x in lines]}

 @r.post("/dispatch/manifests/{manifest_id}/dispatch")
 def dispatch_manifest(manifest_id:int,u=Depends(user_dep)):
  t=tenant(u);results=[]
  with db() as c:
   m=c.execute("SELECT * FROM dist_dispatch_manifests WHERE tenant_id=? AND id=?",(t,manifest_id)).fetchone()
   if not m:raise HTTPException(404,"مانیفست در شرکت شما یافت نشد.")
   if m["status"]!="OPEN":raise HTTPException(409,"فقط مانیفست باز قابل ارسال است.")
   lines=c.execute("SELECT order_id FROM dist_dispatch_manifest_lines WHERE tenant_id=? AND manifest_id=? ORDER BY id",(t,manifest_id)).fetchall()
   if not lines:raise HTTPException(409,"مانیفست خالی قابل ارسال نیست.")
   for line in lines:
    results.append(dispatch_sale(c,t,line["order_id"],m["dispatch_date"],m["manifest_no"],"ارسال گروهی مانیفست "+m["manifest_no"],u.get("id")))
   audit(c,t,"distribution.dispatch_manifest.dispatched:"+m["manifest_no"])
  return {"manifest_id":manifest_id,"manifest_no":m["manifest_no"],"status":"DISPATCHED","invoice_count":len(results),"cost_of_goods":money(sum(minor(x["cost_of_goods"]) for x in results)),"invoices":results}

 @r.post("/dispatch/manifests/{manifest_id}/collections")
 def record_dispatch_collection(manifest_id:int,b:DispatchCollectionIn,u=Depends(user_dep)):
  t=tenant(u)
  try:date.fromisoformat(b.txn_date)
  except (ValueError,TypeError):raise HTTPException(422,"تاریخ دریافت باید YYYY-MM-DD باشد.")
  amount=minor(b.amount,"مبلغ دریافت")
  if amount<=0:raise HTTPException(422,"مبلغ دریافت باید بیشتر از صفر باشد.")
  kind="CASH_COLLECTION" if b.payment_type=="CASH" else "CHECK_COLLECTION"
  cap_col="cash_minor" if b.payment_type=="CASH" else "check_minor"
  gl_code="1010" if b.payment_type=="CASH" else "1020"
  with db() as c:
   m=c.execute("SELECT id,status FROM dist_dispatch_manifests WHERE tenant_id=? AND id=?",(t,manifest_id)).fetchone()
   if not m:raise HTTPException(404,"مانیفست یافت نشد.")
   if m["status"]!="OPEN":raise HTTPException(409,"مانیفست بسته یا لغو شده است.")
   line=c.execute(f"SELECT {cap_col} allocation FROM dist_dispatch_manifest_lines WHERE tenant_id=? AND manifest_id=? AND order_id=?",(t,manifest_id,b.order_id)).fetchone()
   if not line:raise HTTPException(404,"فاکتور متعلق به این مانیفست نیست.")
   prior=c.execute("SELECT COALESCE(SUM(amount_minor),0) FROM dist_dispatch_transactions WHERE tenant_id=? AND manifest_id=? AND order_id=? AND transaction_type=?",(t,manifest_id,b.order_id,kind)).fetchone()[0]
   if prior+amount>line["allocation"]:raise HTTPException(409,"مبلغ دریافت از سهم تعیین‌شده در مانیفست بیشتر است.")
   tid=c.execute("INSERT INTO dist_dispatch_transactions(tenant_id,manifest_id,order_id,transaction_type,amount_minor,txn_date,reference,created_by) VALUES(?,?,?,?,?,?,?,?)",(t,manifest_id,b.order_id,kind,amount,b.txn_date,b.reference.strip(),u.get("id"))).lastrowid
   # Record accepted customer payment for AR aging; account_id=0 denotes funds/checks still held by the distributor, not company treasury.
   c.execute("INSERT INTO dist_customer_receipts(tenant_id,order_id,account_id,txn_date,amount_minor,reference) VALUES(?,?,?,?,?,?)",(t,b.order_id,0,b.txn_date,amount,b.reference.strip() or "DIST-"+str(tid)))
   post(c,t,b.txn_date,"دریافت توزیع‌کننده از فاکتور "+str(b.order_id),kind,tid,[{"code":gl_code,"d":amount},{"code":"1100","c":amount}],u.get("id"))
   audit(c,t,"distribution.dispatch.collection:"+str(tid))
  return {"id":tid,"manifest_id":manifest_id,"order_id":b.order_id,"payment_type":b.payment_type,"amount":money(amount),"status":"RECORDED","journal_posted":True}

 @r.post("/dispatch/manifests/{manifest_id}/remittances")
 def record_dispatch_remittance(manifest_id:int,b:DispatchRemittanceIn,u=Depends(user_dep)):
  t=tenant(u)
  try:date.fromisoformat(b.txn_date)
  except (ValueError,TypeError):raise HTTPException(422,"تاریخ تحویل وجه باید YYYY-MM-DD باشد.")
  amount=minor(b.amount,"مبلغ تحویل وجه")
  if amount<=0:raise HTTPException(422,"مبلغ تحویل وجه باید بیشتر از صفر باشد.")
  with db() as c:
   m=c.execute("SELECT id,status FROM dist_dispatch_manifests WHERE tenant_id=? AND id=?",(t,manifest_id)).fetchone()
   if not m:raise HTTPException(404,"مانیفست یافت نشد.")
   if m["status"]!="OPEN":raise HTTPException(409,"مانیفست بسته یا لغو شده است.")
   account=c.execute("SELECT id,code,name FROM dist_cash_accounts WHERE tenant_id=? AND code=?",(t,b.cash_code.strip())).fetchone()
   if not account:raise HTTPException(404,"حساب صندوق/بانک شرکت یافت نشد.")
   collected=c.execute("SELECT COALESCE(SUM(amount_minor),0) FROM dist_dispatch_transactions WHERE tenant_id=? AND manifest_id=? AND transaction_type='CASH_COLLECTION'",(t,manifest_id)).fetchone()[0]
   remitted=c.execute("SELECT COALESCE(SUM(amount_minor),0) FROM dist_dispatch_transactions WHERE tenant_id=? AND manifest_id=? AND transaction_type='CASH_REMITTANCE'",(t,manifest_id)).fetchone()[0]
   if remitted+amount>collected:raise HTTPException(409,"مبلغ تحویل به شرکت از وجه نقد واقعاً دریافت‌شده بیشتر است.")
   tid=c.execute("INSERT INTO dist_dispatch_transactions(tenant_id,manifest_id,transaction_type,amount_minor,txn_date,reference,account_id,created_by) VALUES(?,?,'CASH_REMITTANCE',?,?,?,?,?)",(t,manifest_id,amount,b.txn_date,b.reference.strip(),account["id"],u.get("id"))).lastrowid
   c.execute("INSERT INTO dist_cash_moves(tenant_id,account_id,txn_date,direction,amount_minor,reference,memo) VALUES(?,?,?,'IN',?,?,?)",(t,account["id"],b.txn_date,amount,b.reference.strip(),"تحویل وجه نقد توزیع‌کننده"))
   post(c,t,b.txn_date,"تحویل وجه نقد توزیع‌کننده","DISTRIBUTOR_REMITTANCE",tid,[{"code":"1000","d":amount},{"code":"1010","c":amount}],u.get("id"))
   audit(c,t,"distribution.dispatch.remittance:"+str(tid))
  return {"id":tid,"manifest_id":manifest_id,"cash_account":account["code"],"amount":money(amount),"status":"RECORDED","journal_posted":True}

 @r.post("/dispatch/transactions/{transaction_id}/handoff-check")
 def handoff_dispatch_check(transaction_id:int,u=Depends(user_dep)):
  t=tenant(u)
  with db() as c:
   tx=c.execute("SELECT * FROM dist_dispatch_transactions WHERE tenant_id=? AND id=? AND transaction_type='CHECK_COLLECTION'",(t,transaction_id)).fetchone()
   if not tx:raise HTTPException(404,"چک دریافت‌شده یافت نشد.")
   if tx["status"]!="RECORDED":raise HTTPException(409,"وضعیت چک اجازه تحویل ندارد.")
   c.execute("UPDATE dist_dispatch_transactions SET status='HANDED_OVER' WHERE tenant_id=? AND id=?",(t,transaction_id))
   audit(c,t,"distribution.dispatch.check_handed_over:"+str(transaction_id))
  return {"id":transaction_id,"status":"HANDED_OVER","note":"تحویل فیزیکی ثبت شد؛ چک هنوز وصول‌شده محسوب نمی‌شود."}

 @r.post("/dispatch/transactions/{transaction_id}/clear-check")
 def clear_dispatch_check(transaction_id:int,payload:dict,u=Depends(user_dep)):
  t=tenant(u)
  try:txn_date=date.fromisoformat(str(payload.get("txn_date") or date.today().isoformat())).isoformat()
  except (ValueError,TypeError):raise HTTPException(422,"تاریخ وصول چک نامعتبر است.")
  with db() as c:
   tx=c.execute("SELECT * FROM dist_dispatch_transactions WHERE tenant_id=? AND id=? AND transaction_type='CHECK_COLLECTION'",(t,transaction_id)).fetchone()
   if not tx:raise HTTPException(404,"چک دریافت‌شده یافت نشد.")
   if tx["status"]!="HANDED_OVER":raise HTTPException(409,"ابتدا تحویل چک به شرکت را ثبت کنید.")
   post(c,t,txn_date,"وصول چک توزیع‌کننده","DISTRIBUTOR_CHECK_CLEARANCE",transaction_id,[{"code":"1000","d":tx["amount_minor"]},{"code":"1020","c":tx["amount_minor"]}],u.get("id"))
   c.execute("UPDATE dist_dispatch_transactions SET status='CLEARED' WHERE tenant_id=? AND id=?",(t,transaction_id))
   audit(c,t,"distribution.dispatch.check_cleared:"+str(transaction_id))
  return {"id":transaction_id,"status":"CLEARED","amount":money(tx["amount_minor"]),"journal_posted":True}

 @r.get("/dispatch/manifests/{manifest_id}/settlement")
 def dispatch_settlement_summary(manifest_id:int,u=Depends(user_dep)):
  t=tenant(u)
  with db() as c:
   m=c.execute("SELECT id,manifest_no,status FROM dist_dispatch_manifests WHERE tenant_id=? AND id=?",(t,manifest_id)).fetchone()
   if not m:raise HTTPException(404,"مانیفست یافت نشد.")
   rows=c.execute("SELECT transaction_type,status,COALESCE(SUM(amount_minor),0) amount FROM dist_dispatch_transactions WHERE tenant_id=? AND manifest_id=? GROUP BY transaction_type,status ORDER BY transaction_type,status",(t,manifest_id)).fetchall()
   txs=c.execute("SELECT id,order_id,transaction_type,amount_minor,txn_date,reference,status FROM dist_dispatch_transactions WHERE tenant_id=? AND manifest_id=? ORDER BY id",(t,manifest_id)).fetchall()
  return {"manifest_id":m["id"],"manifest_no":m["manifest_no"],"manifest_status":m["status"],"totals":[{"type":x["transaction_type"],"status":x["status"],"amount":money(x["amount"])} for x in rows],"transactions":[{"id":x["id"],"order_id":x["order_id"],"type":x["transaction_type"],"amount":money(x["amount_minor"]),"txn_date":x["txn_date"],"reference":x["reference"],"status":x["status"]} for x in txs]}

 @r.get("/sales")
 def sales(u=Depends(user_dep)):
  with db() as c:
   rows=c.execute("SELECT s.*,cu.name customer_name,(SELECT COUNT(*) FROM dist_sale_lines l WHERE l.tenant_id=s.tenant_id AND l.order_id=s.id) line_count FROM dist_sales s JOIN customers cu ON cu.id=s.customer_id AND cu.tenant_id=s.tenant_id WHERE s.tenant_id=? ORDER BY s.id DESC LIMIT 200",(tenant(u),)).fetchall()
  return [{"id":x["id"],"order_no":x["order_no"],"order_date":x["order_date"],"status":x["status"],"total":money(x["total_minor"]),"customer_name":x["customer_name"],"line_count":x["line_count"]} for x in rows]
 @r.get("/sales/{oid}")
 def sale_detail(oid:int,u=Depends(user_dep)):
  t=tenant(u)
  with db() as c:
   order=c.execute("SELECT s.*,cu.name customer_name FROM dist_sales s JOIN customers cu ON cu.id=s.customer_id AND cu.tenant_id=s.tenant_id WHERE s.tenant_id=? AND s.id=?",(t,oid)).fetchone()
   if not order:raise HTTPException(404,"فاکتور در شرکت شما یافت نشد.")
   lines=c.execute("SELECT l.product_id,p.sku,p.name,l.quantity,l.unit_price_minor,l.discount_minor,l.total_minor FROM dist_sale_lines l JOIN products p ON p.id=l.product_id AND p.tenant_id=l.tenant_id WHERE l.tenant_id=? AND l.order_id=? ORDER BY l.id",(t,oid)).fetchall()
  return {"id":order["id"],"order_no":order["order_no"],"customer_id":order["customer_id"],"customer_name":order["customer_name"],"order_date":order["order_date"],"status":order["status"],"total":money(order["total_minor"]),"lines":[{"product_id":x["product_id"],"sku":x["sku"],"name":x["name"],"quantity":x["quantity"],"unit_price":money(x["unit_price_minor"]),"discount":money(x["discount_minor"]),"total":money(x["total_minor"])} for x in lines]}
 @r.post("/sales")
 def new_sale(b:Sale,u=Depends(user_dep)):
  t=tenant(u); total=0; lines=[]
  try:date.fromisoformat(b.order_date)
  except ValueError:raise HTTPException(422,"تاریخ سفارش نامعتبر است.")
  with db() as c:
   if not c.execute("SELECT 1 FROM customers WHERE tenant_id=? AND id=?",(t,b.customer_id)).fetchone():raise HTTPException(404,"مشتری در شرکت شما یافت نشد.")
   for x in b.lines:
    if not c.execute("SELECT 1 FROM products WHERE tenant_id=? AND id=?",(t,x.product_id)).fetchone():raise HTTPException(404,"کالا در شرکت شما یافت نشد.")
    gross=minor(x.unit_price*Decimal(str(x.quantity))); discount=minor(x.discount)
    if discount>gross:raise HTTPException(422,"تخفیف از مبلغ ردیف بیشتر است.")
    lt=gross-discount; total+=lt; lines.append((x.product_id,x.quantity,minor(x.unit_price),discount,lt))
   if total<=0:raise HTTPException(422,"مبلغ سفارش باید بیشتر از صفر باشد.")
   try:oid=c.execute("INSERT INTO dist_sales(tenant_id,order_no,customer_id,order_date,total_minor,note) VALUES(?,?,?,?,?,?)",(t,b.order_no,b.customer_id,b.order_date,total,b.note)).lastrowid
   except sqlite3.IntegrityError:raise HTTPException(409,"شماره سفارش تکراری است.")
   c.executemany("INSERT INTO dist_sale_lines(tenant_id,order_id,product_id,quantity,unit_price_minor,discount_minor,total_minor) VALUES(?,?,?,?,?,?,?)",[(t,oid,*x) for x in lines]);audit(c,t,"distribution.sale.draft:"+b.order_no)
  return {"id":oid,"order_no":b.order_no,"status":"DRAFT","total":money(total)}
 @r.post("/sales/{oid}/post")
 def post_sale(oid:int,u=Depends(user_dep)):
  t=tenant(u)
  with db() as c:
   o=c.execute("SELECT * FROM dist_sales WHERE tenant_id=? AND id=?",(t,oid)).fetchone()
   if not o:raise HTTPException(404,"سفارش یافت نشد.")
   if o["status"]!="DRAFT":raise HTTPException(409,"فقط سفارش پیش‌نویس قابل نهایی‌سازی است.")
   lines=c.execute("SELECT l.*,COALESCE(d.cost_minor,0) cost,COALESCE(d.warehouse_code,'DEFAULT') wh FROM dist_sale_lines l JOIN products p ON p.id=l.product_id AND p.tenant_id=l.tenant_id LEFT JOIN dist_item_data d ON d.tenant_id=p.tenant_id AND d.product_id=p.id WHERE l.tenant_id=? AND l.order_id=?",(t,oid)).fetchall()
   revenue=sum(x["total_minor"] for x in lines); cost=sum(minor(Decimal(str(x["quantity"]))*Decimal(x["cost"])/100) for x in lines)
   # Invoice posting recognizes revenue/receivable only. Inventory and COGS move at physical dispatch.
   post(c,t,o["order_date"],"فروش "+o["order_no"],"SALE",oid,[{"code":"1100","d":revenue},{"code":"4000","c":revenue}],u.get("id"))
   c.execute("UPDATE dist_sales SET status='POSTED' WHERE tenant_id=? AND id=?",(t,oid));audit(c,t,"distribution.sale.posted:"+o["order_no"])
  return {"status":"POSTED","revenue":money(revenue),"cost_of_goods":"0.00","pending_dispatch_cost":money(cost),"gross_profit":None}

 @r.post("/sales/{oid}/deliver")
 def deliver(oid:int,payload:dict,u=Depends(user_dep)):
  t=tenant(u)
  dispatch_date=str(payload.get("delivery_date") or date.today().isoformat())
  result={}
  with db() as c:
   result=dispatch_sale(c,t,oid,dispatch_date,str(payload.get("reference",""))[:120],str(payload.get("note",""))[:500],u.get("id"))
  return result
 @r.post("/purchases")
 def new_purchase(b:Purchase,u=Depends(user_dep)):
  t=tenant(u);total=0;lines=[]
  with db() as c:
   if not c.execute("SELECT 1 FROM dist_suppliers WHERE tenant_id=? AND id=?",(t,b.supplier_id)).fetchone():raise HTTPException(404,"تأمین‌کننده یافت نشد.")
   if not c.execute("SELECT 1 FROM dist_warehouses WHERE tenant_id=? AND code=?",(t,b.warehouse_code)).fetchone():raise HTTPException(404,"انبار یافت نشد.")
   for x in b.lines:
    if not c.execute("SELECT 1 FROM products WHERE tenant_id=? AND id=?",(t,x.product_id)).fetchone():raise HTTPException(404,"کالا یافت نشد.")
    cost=minor(x.unit_cost);lt=minor(x.unit_cost*Decimal(str(x.quantity)));total+=lt;lines.append((x.product_id,x.quantity,cost,lt))
   if total<=0:raise HTTPException(422,"مبلغ خرید باید بیشتر از صفر باشد.")
   try:oid=c.execute("INSERT INTO dist_purchases(tenant_id,order_no,supplier_id,warehouse_code,order_date,total_minor) VALUES(?,?,?,?,?,?)",(t,b.order_no,b.supplier_id,b.warehouse_code,b.order_date,total)).lastrowid
   except sqlite3.IntegrityError:raise HTTPException(409,"شماره خرید تکراری است.")
   c.executemany("INSERT INTO dist_purchase_lines(tenant_id,order_id,product_id,quantity,cost_minor,total_minor) VALUES(?,?,?,?,?,?)",[(t,oid,*x) for x in lines]);audit(c,t,"distribution.purchase.draft:"+b.order_no)
  return {"id":oid,"order_no":b.order_no,"status":"DRAFT","total":money(total)}
 @r.post("/purchases/{oid}/receive")
 def receive_purchase(oid:int,u=Depends(user_dep)):
  t=tenant(u)
  with db() as c:
   o=c.execute("SELECT * FROM dist_purchases WHERE tenant_id=? AND id=?",(t,oid)).fetchone()
   if not o:raise HTTPException(404,"سفارش خرید یافت نشد.")
   if o["status"]!="DRAFT":raise HTTPException(409,"فقط خرید پیش‌نویس قابل دریافت است.")
   lines=c.execute("SELECT * FROM dist_purchase_lines WHERE tenant_id=? AND order_id=?",(t,oid)).fetchall();total=sum(x["total_minor"] for x in lines)
   post(c,t,o["order_date"],"خرید "+o["order_no"],"PURCHASE",oid,[{"code":"1200","d":total},{"code":"2000","c":total}],u.get("id"))
   for x in lines:
    c.execute("UPDATE products SET stock=stock+? WHERE tenant_id=? AND id=?",(x["quantity"],t,x["product_id"]))
    c.execute("INSERT INTO dist_item_data(tenant_id,product_id,cost_minor,warehouse_code) VALUES(?,?,?,?) ON CONFLICT(tenant_id,product_id) DO UPDATE SET cost_minor=excluded.cost_minor,warehouse_code=excluded.warehouse_code",(t,x["product_id"],x["cost_minor"],o["warehouse_code"]))
    c.execute("INSERT INTO dist_stock_moves(tenant_id,product_id,warehouse_code,quantity,cost_minor,kind,reference,created_by) VALUES(?,?,?,?,?,?,?,?)",(t,x["product_id"],o["warehouse_code"],x["quantity"],x["cost_minor"],"PURCHASE_IN",o["order_no"],u.get("id")))
   c.execute("UPDATE dist_purchases SET status='RECEIVED' WHERE tenant_id=? AND id=?",(t,oid));audit(c,t,"distribution.purchase.received:"+o["order_no"])
  return {"status":"RECEIVED","inventory_value":money(total)}
 @r.post("/accounts")
 def add_account(b:AccountIn,u=Depends(user_dep)):
  t=tenant(u)
  with db() as c:
   if b.parent_code and not c.execute("SELECT 1 FROM dist_accounts WHERE tenant_id=? AND code=?",(t,b.parent_code)).fetchone():raise HTTPException(422,"حساب والد یافت نشد.")
   try:c.execute("INSERT INTO dist_accounts(tenant_id,code,name,nature,parent_code) VALUES(?,?,?,?,?)",(t,b.code.strip(),b.name.strip(),b.nature,b.parent_code))
   except sqlite3.IntegrityError:raise HTTPException(409,"کد حساب تکراری است.")
   audit(c,t,"distribution.account.created:"+b.code)
  return {"status":"created","code":b.code}
 @r.post("/journals")
 def manual_journal(b:JournalIn,u=Depends(user_dep)):
  t=tenant(u);lines=[]
  for x in b.lines:lines.append({"code":x.account_code,"d":minor(x.debit),"c":minor(x.credit)})
  with db() as c:
   no=post(c,t,b.entry_date,b.description,"MANUAL",None,lines,u.get("id"));audit(c,t,"distribution.journal.posted:"+no)
  return {"status":"POSTED","entry_no":no}
 @r.get("/journals")
 def journal_list(limit:int=Query(100,ge=1,le=1000),u=Depends(user_dep)):
  with db() as c:
   rows=c.execute("SELECT j.id,j.entry_no,j.entry_date,j.description,j.source_type,COALESCE(SUM(l.debit_minor),0) d,COALESCE(SUM(l.credit_minor),0) cr FROM dist_journals j LEFT JOIN dist_journal_lines l ON l.journal_id=j.id AND l.tenant_id=j.tenant_id WHERE j.tenant_id=? GROUP BY j.id ORDER BY j.id DESC LIMIT ?",(tenant(u),limit)).fetchall()
  return [{"id":x["id"],"entry_no":x["entry_no"],"entry_date":x["entry_date"],"description":x["description"],"source_type":x["source_type"],"debit":money(x["d"]),"credit":money(x["cr"])} for x in rows]
 @r.get("/reports/ledger/{code}")
 def ledger(code:str,u=Depends(user_dep)):
  t=tenant(u)
  with db() as c:
   if not c.execute("SELECT 1 FROM dist_accounts WHERE tenant_id=? AND code=?",(t,code)).fetchone():raise HTTPException(404,"حساب یافت نشد.")
   rows=c.execute("SELECT j.entry_no,j.entry_date,j.description,l.debit_minor,l.credit_minor FROM dist_journal_lines l JOIN dist_journals j ON j.id=l.journal_id AND j.tenant_id=l.tenant_id WHERE l.tenant_id=? AND l.account_code=? ORDER BY j.entry_date,j.id,l.id",(t,code)).fetchall()
  balance=0;items=[]
  for x in rows:
   balance+=x["debit_minor"]-x["credit_minor"];items.append({"entry_no":x["entry_no"],"date":x["entry_date"],"description":x["description"],"debit":money(x["debit_minor"]),"credit":money(x["credit_minor"]),"balance":money(balance)})
  return {"account_code":code,"items":items,"balance":money(balance)}
 @r.get("/reports/balance-sheet")
 def balance_sheet(u=Depends(user_dep)):
  with db() as c:
   rows=c.execute("SELECT a.nature,COALESCE(SUM(l.debit_minor-l.credit_minor),0) bal FROM dist_accounts a LEFT JOIN dist_journal_lines l ON l.tenant_id=a.tenant_id AND l.account_code=a.code WHERE a.tenant_id=? AND a.nature IN ('asset','liability','equity','income','expense') GROUP BY a.nature",(tenant(u),)).fetchall()
  vals={x["nature"]:x["bal"] for x in rows};a=vals.get("asset",0);liab=-vals.get("liability",0);eq=-vals.get("equity",0);profit=-(vals.get("income",0))-vals.get("expense",0);eq+=profit
  return {"assets":money(a),"liabilities":money(liab),"equity_including_current_profit":money(eq),"current_profit":money(profit),"balanced":a-liab-eq==0}
 @r.post("/sales/receipt")
 def sale_receipt(b:ReceiptIn,u=Depends(user_dep)):
  t=tenant(u);a=minor(b.amount)
  with db() as c:
   o=c.execute("SELECT * FROM dist_sales WHERE tenant_id=? AND id=?",(t,b.order_id)).fetchone()
   if not o or o["status"] not in ("POSTED","DELIVERED"):raise HTTPException(404,"فاکتور نهایی‌شده یافت نشد.")
   acc=c.execute("SELECT * FROM dist_cash_accounts WHERE tenant_id=? AND code=?",(t,b.cash_code)).fetchone()
   if not acc:raise HTTPException(404,"حساب صندوق/بانک یافت نشد.")
   received=c.execute("SELECT COALESCE(SUM(amount_minor),0) FROM dist_customer_receipts WHERE tenant_id=? AND order_id=?",(t,b.order_id)).fetchone()[0]
   if a>o["total_minor"]-received:raise HTTPException(409,"دریافت از مانده فاکتور بیشتر است.")
   try:date.fromisoformat(b.txn_date)
   except ValueError:raise HTTPException(422,"تاریخ دریافت نامعتبر است.")
   c.execute("INSERT INTO dist_customer_receipts(tenant_id,order_id,account_id,txn_date,amount_minor,reference) VALUES(?,?,?,?,?,?)",(t,b.order_id,acc["id"],b.txn_date,a,b.reference))
   c.execute("INSERT INTO dist_cash_moves(tenant_id,account_id,txn_date,direction,amount_minor,reference,memo) VALUES(?,?,?,?,?,?,?)",(t,acc["id"],b.txn_date,"IN",a,b.reference,"وصول فاکتور "+o["order_no"]))
   post(c,t,b.txn_date,"وصول فاکتور "+o["order_no"],"RECEIPT",b.order_id,[{"code":"1000","d":a},{"code":"1100","c":a}],u.get("id"));audit(c,t,"distribution.receipt:"+o["order_no"])
  return {"status":"POSTED","received":money(a),"remaining":money(o["total_minor"]-received-a)}
 @r.get("/returns")
 def list_returns(u=Depends(user_dep)):
  t=tenant(u)
  with db() as c:
   rows=c.execute("SELECT r.*,s.order_no FROM dist_sales_returns r JOIN dist_sales s ON s.id=r.order_id AND s.tenant_id=r.tenant_id WHERE r.tenant_id=? ORDER BY r.id DESC LIMIT 200",(t,)).fetchall()
   out=[]
   for r0 in rows:
    lines=c.execute("SELECT product_id,quantity,revenue_minor,cost_minor FROM dist_sales_return_lines WHERE tenant_id=? AND return_id=? ORDER BY id",(t,r0["id"])).fetchall()
    out.append({"id":r0["id"],"return_no":r0["return_no"],"order_id":r0["order_id"],"order_no":r0["order_no"],"return_date":r0["return_date"],"reason":r0["reason"],"status":r0["status"],"lines":[{"product_id":x["product_id"],"quantity":x["quantity"],"revenue":money(x["revenue_minor"]),"cost":money(x["cost_minor"])} for x in lines]})
  return out

 @r.post("/returns")
 def create_return(b:SalesReturnIn,u=Depends(user_dep)):
  t=tenant(u)
  try:date.fromisoformat(b.return_date)
  except ValueError:raise HTTPException(422,"تاریخ برگشت نامعتبر است.")
  if len({x.product_id for x in b.lines})!=len(b.lines):raise HTTPException(422,"هر کالا باید فقط یک بار در ردیف‌های برگشت درج شود.")
  with db() as c:
   order=c.execute("SELECT * FROM dist_sales WHERE tenant_id=? AND id=?",(t,b.order_id)).fetchone()
   if not order:raise HTTPException(404,"فاکتور در شرکت شما یافت نشد.")
   if order["status"]!="DELIVERED":raise HTTPException(409,"فقط کالای ارسال‌شده قابل برگشت است.")
   sold=c.execute("SELECT product_id,SUM(quantity) quantity FROM dist_sale_lines WHERE tenant_id=? AND order_id=? GROUP BY product_id",(t,b.order_id)).fetchall()
   sold_by={x["product_id"]:Decimal(str(x["quantity"])) for x in sold}
   for line in b.lines:
    if line.product_id not in sold_by:raise HTTPException(422,"کالای برگشتی در فاکتور اصلی وجود ندارد.")
    previous=c.execute("SELECT COALESCE(SUM(l.quantity),0) FROM dist_sales_return_lines l JOIN dist_sales_returns r ON r.id=l.return_id AND r.tenant_id=l.tenant_id WHERE l.tenant_id=? AND r.order_id=? AND l.product_id=? AND r.status='APPROVED'",(t,b.order_id,line.product_id)).fetchone()[0]
    if Decimal(str(previous))+line.quantity>sold_by[line.product_id]:raise HTTPException(409,"مقدار برگشتی از مقدار فروخته‌شده بیشتر است.")
   try:
    rid=c.execute("INSERT INTO dist_sales_returns(tenant_id,order_id,return_no,return_date,reason,created_by) VALUES(?,?,?,?,?,?)",(t,b.order_id,b.return_no.strip(),b.return_date,b.reason.strip(),u.get("id"))).lastrowid
    c.executemany("INSERT INTO dist_sales_return_lines(tenant_id,return_id,product_id,quantity) VALUES(?,?,?,?)",[(t,rid,x.product_id,float(x.quantity)) for x in b.lines])
   except sqlite3.IntegrityError:raise HTTPException(409,"شماره برگشت تکراری است یا ردیف برگشت نامعتبر است.")
   audit(c,t,"distribution.return.draft:"+b.return_no.strip())
  return {"id":rid,"return_no":b.return_no.strip(),"status":"DRAFT"}

 @r.post("/returns/{return_id}/approve")
 def approve_return(return_id:int,u=Depends(user_dep)):
  t=tenant(u)
  with db() as c:
   ret=c.execute("SELECT r.*,s.order_no,s.status order_status FROM dist_sales_returns r JOIN dist_sales s ON s.id=r.order_id AND s.tenant_id=r.tenant_id WHERE r.tenant_id=? AND r.id=?",(t,return_id)).fetchone()
   if not ret:raise HTTPException(404,"برگشت کالا یافت نشد.")
   if ret["status"]!="DRAFT":raise HTTPException(409,"فقط برگشت پیش‌نویس قابل تأیید است.")
   if ret["order_status"]!="DELIVERED":raise HTTPException(409,"فقط کالای ارسال‌شده قابل تأیید برگشت است.")
   lines=c.execute("SELECT * FROM dist_sales_return_lines WHERE tenant_id=? AND return_id=? ORDER BY id",(t,return_id)).fetchall()
   revenue_total=0;cost_total=0;stock_updates=[]
   for line in lines:
    sale=c.execute("SELECT COALESCE(SUM(quantity),0) qty,COALESCE(SUM(total_minor),0) revenue FROM dist_sale_lines WHERE tenant_id=? AND order_id=? AND product_id=?",(t,ret["order_id"],line["product_id"])).fetchone()
    if not sale or Decimal(str(sale["qty"]))<=0:raise HTTPException(409,"ردیف فاکتور اصلی یافت نشد.")
    previous=c.execute("SELECT COALESCE(SUM(l.quantity),0) FROM dist_sales_return_lines l JOIN dist_sales_returns r ON r.id=l.return_id AND r.tenant_id=l.tenant_id WHERE l.tenant_id=? AND r.order_id=? AND l.product_id=? AND r.status='APPROVED'",(t,ret["order_id"],line["product_id"])).fetchone()[0]
    if Decimal(str(previous))+Decimal(str(line["quantity"]))>Decimal(str(sale["qty"])):raise HTTPException(409,"مقدار برگشتی تأییدشده از مقدار فروش بیشتر می‌شود.")
    revenue=minor(Decimal(str(sale["revenue"]))*Decimal(str(line["quantity"]))/Decimal(str(sale["qty"]))/100)
    unit_cost=c.execute("SELECT cost_minor FROM dist_stock_moves WHERE tenant_id=? AND product_id=? AND reference=? AND kind='SALE_OUT' ORDER BY id LIMIT 1",(t,line["product_id"],ret["order_no"])).fetchone()
    cost=minor(Decimal(str(unit_cost["cost_minor"] if unit_cost else 0))*Decimal(str(line["quantity"]))/100)
    product=c.execute("SELECT stock FROM products WHERE tenant_id=? AND id=?",(t,line["product_id"])).fetchone()
    if not product:raise HTTPException(404,"کالا در شرکت شما یافت نشد.")
    revenue_total+=revenue;cost_total+=cost
    stock_updates.append((line["product_id"],float(line["quantity"]),revenue,cost))
   if revenue_total<=0:raise HTTPException(422,"مبلغ برگشت باید بیشتر از صفر باشد.")
   post(c,t,ret["return_date"],"برگشت فروش "+ret["return_no"],"SALES_RETURN",return_id,[{"code":"4000","d":revenue_total},{"code":"1100","c":revenue_total}],u.get("id"))
   if cost_total:
    post(c,t,ret["return_date"],"برگشت بهای تمام‌شده "+ret["return_no"],"RETURN_COGS",return_id,[{"code":"1200","d":cost_total},{"code":"5000","c":cost_total}],u.get("id"))
   for product_id,quantity,revenue,cost in stock_updates:
    c.execute("UPDATE products SET stock=stock+? WHERE tenant_id=? AND id=?",(quantity,t,product_id))
    wh=c.execute("SELECT warehouse_code FROM dist_item_data WHERE tenant_id=? AND product_id=?",(t,product_id)).fetchone()
    c.execute("INSERT INTO dist_stock_moves(tenant_id,product_id,warehouse_code,quantity,cost_minor,kind,reference,created_by) VALUES(?,?,?,?,?,?,?,?)",(t,product_id,wh["warehouse_code"] if wh else "DEFAULT",quantity,minor(Decimal(cost)/100/Decimal(str(quantity))) if quantity else 0,"RETURN_IN",ret["return_no"],u.get("id")))
    c.execute("UPDATE dist_sales_return_lines SET revenue_minor=?,cost_minor=? WHERE tenant_id=? AND return_id=? AND product_id=?",(revenue,cost,t,return_id,product_id))
   c.execute("UPDATE dist_sales_returns SET status='APPROVED',approved_at=CURRENT_TIMESTAMP WHERE tenant_id=? AND id=? AND status='DRAFT'",(t,return_id))
   audit(c,t,"distribution.return.approved:"+ret["return_no"])
  return {"status":"APPROVED","return_no":ret["return_no"],"revenue_reversed":money(revenue_total),"cost_reversed":money(cost_total),"inventory_restored":True,"settlement":"CUSTOMER_CREDIT_NOTE"}

 @r.get("/reports/receivables")
 def receivables(u=Depends(user_dep)):
  t=tenant(u)
  with db() as c:
   opening=c.execute("SELECT code,name,SUM(amount_minor) amount FROM dist_opening_balances WHERE tenant_id=? AND kind='CUSTOMER' GROUP BY code,name",(t,)).fetchall()
   rows=c.execute("SELECT cu.id,cu.name,COALESCE(cc.code,'#'||cu.id) code,COALESCE(SUM(s.total_minor),0)-COALESCE((SELECT SUM(rl.revenue_minor) FROM dist_sales_returns r JOIN dist_sales_return_lines rl ON rl.return_id=r.id AND rl.tenant_id=r.tenant_id JOIN dist_sales rs ON rs.id=r.order_id AND rs.tenant_id=r.tenant_id WHERE r.tenant_id=cu.tenant_id AND rs.customer_id=cu.id AND r.status='APPROVED'),0) total FROM customers cu LEFT JOIN dist_customer_codes cc ON cc.customer_id=cu.id AND cc.tenant_id=cu.tenant_id LEFT JOIN dist_sales s ON s.customer_id=cu.id AND s.tenant_id=cu.tenant_id AND s.status IN ('POSTED','DELIVERED') WHERE cu.tenant_id=? GROUP BY cu.id,cu.name,cc.code",(t,)).fetchall()
   rr=c.execute("SELECT s.customer_id,SUM(r.amount_minor) amount FROM dist_customer_receipts r JOIN dist_sales s ON s.id=r.order_id AND s.tenant_id=r.tenant_id WHERE r.tenant_id=? GROUP BY s.customer_id",(t,)).fetchall()
  paid={x["customer_id"]:x["amount"] for x in rr};items=[{"code":x["code"],"name":x["name"],"opening":money(x["amount"]),"sales":"0.00","received":"0.00","outstanding":money(x["amount"])} for x in opening]
  for x in rows:
   if x["total"] or x["id"] in paid:items.append({"code":x["code"],"name":x["name"],"opening":"0.00","sales":money(x["total"]),"received":money(paid.get(x["id"],0)),"outstanding":money(x["total"]-paid.get(x["id"],0))})
  return {"items":items,"total_outstanding":money(sum(minor(x["outstanding"]) for x in items))}
 @r.get("/reports/payables")
 def payables(u=Depends(user_dep)):
  t=tenant(u)
  with db() as c:
   opening=c.execute("SELECT code,name,SUM(amount_minor) amount FROM dist_opening_balances WHERE tenant_id=? AND kind='SUPPLIER' GROUP BY code,name",(t,)).fetchall()
   rows=c.execute("SELECT s.code,s.name,COALESCE(SUM(p.total_minor),0) total FROM dist_suppliers s LEFT JOIN dist_purchases p ON p.supplier_id=s.id AND p.tenant_id=s.tenant_id AND p.status='RECEIVED' WHERE s.tenant_id=? GROUP BY s.id",(t,)).fetchall()
  items=[{"code":x["code"],"name":x["name"],"opening":money(x["amount"]),"purchases":"0.00","outstanding":money(x["amount"])} for x in opening]
  items.extend({"code":x["code"],"name":x["name"],"opening":"0.00","purchases":money(x["total"]),"outstanding":money(x["total"])} for x in rows if x["total"])
  return {"items":items,"total_outstanding":money(sum(minor(x["outstanding"]) for x in items))}
 @r.get("/reports/trial-balance")
 def trial(u=Depends(user_dep)):
  with db() as c:
   rows=c.execute("SELECT a.code,a.name,a.nature,COALESCE(SUM(l.debit_minor),0) d,COALESCE(SUM(l.credit_minor),0) cr FROM dist_accounts a LEFT JOIN dist_journal_lines l ON l.tenant_id=a.tenant_id AND l.account_code=a.code LEFT JOIN dist_journals j ON j.id=l.journal_id AND j.tenant_id=a.tenant_id AND j.status='POSTED' WHERE a.tenant_id=? GROUP BY a.id ORDER BY a.code",(tenant(u),)).fetchall()
  return {"items":[{"code":x["code"],"name":x["name"],"nature":x["nature"],"debit":money(x["d"]),"credit":money(x["cr"]),"balance":money(x["d"]-x["cr"])} for x in rows],"balanced":sum(x["d"] for x in rows)==sum(x["cr"] for x in rows)}
 @r.get("/reports/profit-loss")
 def pl(u=Depends(user_dep)):
  with db() as c:
   rows=c.execute("SELECT a.nature,COALESCE(SUM(l.debit_minor),0) d,COALESCE(SUM(l.credit_minor),0) cr FROM dist_accounts a LEFT JOIN dist_journal_lines l ON l.tenant_id=a.tenant_id AND l.account_code=a.code WHERE a.tenant_id=? AND a.nature IN ('income','expense') GROUP BY a.nature",(tenant(u),)).fetchall()
  m={x["nature"]:x for x in rows};inc=m["income"]["cr"]-m["income"]["d"] if "income" in m else 0;exp=m["expense"]["d"]-m["expense"]["cr"] if "expense" in m else 0
  return {"income":money(inc),"expenses":money(exp),"profit":money(inc-exp)}
 @r.get("/reports/stock-movements")
 def stock_moves(limit:int=Query(100,ge=1,le=1000),u=Depends(user_dep)):
  with db() as c:return [dict(x) for x in c.execute("SELECT m.*,p.sku,p.name FROM dist_stock_moves m JOIN products p ON p.id=m.product_id AND p.tenant_id=m.tenant_id WHERE m.tenant_id=? ORDER BY m.id DESC LIMIT ?",(tenant(u),limit)).fetchall()]
 @r.get("/treasury/accounts")
 def cash_accounts(u=Depends(user_dep)):
  with db() as c:
   rows=c.execute("SELECT a.code,a.name,a.kind,a.opening_minor+COALESCE(SUM(CASE WHEN m.direction='IN' THEN m.amount_minor ELSE -m.amount_minor END),0) bal FROM dist_cash_accounts a LEFT JOIN dist_cash_moves m ON m.account_id=a.id AND m.tenant_id=a.tenant_id WHERE a.tenant_id=? GROUP BY a.id ORDER BY a.code",(tenant(u),)).fetchall()
  return [{"code":x["code"],"name":x["name"],"kind":x["kind"],"balance":money(x["bal"])} for x in rows]
 @r.post("/treasury/accounts")
 def add_cash(b:CashAccount,u=Depends(user_dep)):
  t=tenant(u);opening=minor(b.opening_balance)
  with db() as c:
   try:aid=c.execute("INSERT INTO dist_cash_accounts(tenant_id,code,name,kind,opening_minor) VALUES(?,?,?,?,?)",(t,b.code,b.name,b.kind,opening)).lastrowid
   except sqlite3.IntegrityError:raise HTTPException(409,"کد صندوق/بانک تکراری است.")
   if opening:post(c,t,date.today().isoformat(),"مانده افتتاحیه "+b.code,"CASH_OPENING",aid,[{"code":"1000","d":opening},{"code":"3000","c":opening}],u.get("id"))
   audit(c,t,"distribution.cash.created:"+b.code)
  return {"id":aid,"code":b.code,"opening_balance":money(opening)}
 @r.post("/treasury/movements")
 def cash_move(b:CashMove,u=Depends(user_dep)):
  t=tenant(u);a=minor(b.amount)
  try:date.fromisoformat(b.txn_date)
  except ValueError:raise HTTPException(422,"تاریخ نامعتبر است.")
  with db() as c:
   acc=c.execute("SELECT * FROM dist_cash_accounts WHERE tenant_id=? AND code=?",(t,b.code)).fetchone()
   if not acc:raise HTTPException(404,"صندوق/بانک یافت نشد.")
   bal=acc["opening_minor"]+c.execute("SELECT COALESCE(SUM(CASE WHEN direction='IN' THEN amount_minor ELSE -amount_minor END),0) FROM dist_cash_moves WHERE tenant_id=? AND account_id=?",(t,acc["id"])).fetchone()[0]
   if b.direction=="OUT" and bal<a:raise HTTPException(409,"مانده برای پرداخت کافی نیست.")
   c.execute("INSERT INTO dist_cash_moves(tenant_id,account_id,txn_date,direction,amount_minor,reference,memo) VALUES(?,?,?,?,?,?,?)",(t,acc["id"],b.txn_date,b.direction,a,b.reference,b.memo))
   lines=[{"code":"1000","d":a},{"code":"1100","c":a}] if b.direction=="IN" else [{"code":"2000","d":a},{"code":"1000","c":a}]
   post(c,t,b.txn_date,b.memo or "تراکنش خزانه","TREASURY",None,lines,u.get("id"));audit(c,t,"distribution.treasury."+b.direction)
  return {"status":"POSTED","amount":money(a)}
 @r.get("/import/catalog")
 def import_catalog(u=Depends(user_dep)):
  return [{"key":k,"title":v["title"],"columns":v["columns"],"required":v["required"]} for k,v in SCHEMAS.items()]
 @r.get("/import/template/{dataset}")
 def template(dataset:str,u=Depends(user_dep)):
  if dataset not in SCHEMAS:raise HTTPException(404,"الگوی ناشناخته است.")
  try:
   from openpyxl import Workbook
   w=Workbook();s=w.active;s.title=SCHEMAS[dataset]["title"][:31];s.append(SCHEMAS[dataset]["columns"]);s.freeze_panes="A2";s.auto_filter.ref=s.dimensions
   for col in s.columns:s.column_dimensions[col[0].column_letter].width=min(max(len(str(col[0].value))+4,16),32)
   out=io.BytesIO();w.save(out);out.seek(0)
  except Exception as e:raise HTTPException(500,"ساخت الگو ناموفق بود.") from e
  return StreamingResponse(out,media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",headers={"Content-Disposition":f'attachment; filename="leadora_{dataset}_template.xlsx"',"Cache-Control":"no-store"})
 @r.post("/import/{dataset}/preview")
 async def preview(dataset:str,file:UploadFile=File(...),u=Depends(user_dep)):
  if dataset not in SCHEMAS:raise HTTPException(404,"نوع داده ناشناخته است.")
  raw=await file.read(LIMIT_BYTES+1)
  if len(raw)>LIMIT_BYTES:raise HTTPException(413,"حداکثر حجم فایل ۱۰ مگابایت است.")
  if not file.filename or not file.filename.lower().endswith(".xlsx"):raise HTTPException(415,"فقط فایل .xlsx پشتیبانی می‌شود.")
  try:
   from openpyxl import load_workbook
   w=load_workbook(io.BytesIO(raw),read_only=True,data_only=True);it=w.active.iter_rows(values_only=True);headers=[str(x).strip() if x is not None else "" for x in next(it,())]
   missing=[x for x in SCHEMAS[dataset]["required"] if x not in headers]
   if missing:raise HTTPException(422,"ستون‌های ضروری موجود نیست: "+"، ".join(missing))
   rows=[];errors=[]
   for n,vals in enumerate(it,2):
    if n>10001:raise HTTPException(413,"حداکثر ۱۰ هزار ردیف مجاز است.")
    row={headers[i]:vals[i] for i in range(min(len(headers),len(vals))) if headers[i]}
    if not any(v not in (None,"") for v in row.values()):continue
    for key in SCHEMAS[dataset]["required"]:
     if row.get(key) in (None,""):errors.append({"row":n,"field":key,"error":"مقدار ضروری خالی است."})
    fields={"items":["قیمت خرید","قیمت فروش","موجودی اولیه","حداقل موجودی"],"banks":["مانده افتتاحیه"],"cash":["مانده افتتاحیه"],"receivables":["مبلغ طلب"],"payables":["مبلغ بدهی"],"opening_finance":["بدهکار افتتاحیه","بستانکار افتتاحیه"],"customers":["مانده طلب افتتاحیه"],"suppliers":["مانده بدهی افتتاحیه"]}.get(dataset,[])
    for field in fields:
     try:
      if field in ("موجودی اولیه","حداقل موجودی"):
       if float(row.get(field) or 0)<0:raise ValueError()
      elif minor(row.get(field,0),field)<0:raise ValueError()
     except (ValueError,TypeError,HTTPException):errors.append({"row":n,"field":field,"error":"مقدار باید عددی و نامنفی باشد."})
    if dataset=="opening_finance":
     d=minor(row.get("بدهکار افتتاحیه",0));cr=minor(row.get("بستانکار افتتاحیه",0))
     if d and cr:errors.append({"row":n,"field":"بدهکار/بستانکار","error":"در یک ردیف فقط یکی از بدهکار یا بستانکار مجاز است."})
    rows.append(row)
   w.close()
   if dataset=="opening_finance" and not errors:
    d=sum(minor(x.get("بدهکار افتتاحیه")) for x in rows);cr=sum(minor(x.get("بستانکار افتتاحیه")) for x in rows)
    if d!=cr:errors.append({"row":0,"error":"جمع بدهکار و بستانکار افتتاحیه برابر نیست."})
  except HTTPException:raise
  except Exception as e:raise HTTPException(422,"فایل اکسل قابل خواندن نیست.") from e
  t=tenant(u);bid=secrets.token_urlsafe(16);report={"dataset":dataset,"total_rows":len(rows),"errors":errors[:500],"valid":not errors,"status":"PREVIEW"}
  with db() as c:c.execute("INSERT INTO dist_import_batches(id,tenant_id,dataset,rows_json,report_json) VALUES(?,?,?,?,?)",(bid,t,dataset,json.dumps(rows,ensure_ascii=False,default=str),json.dumps(report,ensure_ascii=False)))
  return {"batch_id":bid,**report,"sample":rows[:10]}
 @r.post("/import/{bid}/commit")
 def commit_import(bid:str,u=Depends(user_dep)):
  t=tenant(u)
  with db() as c:
   b=c.execute("SELECT * FROM dist_import_batches WHERE id=? AND tenant_id=?",(bid,t)).fetchone()
   if not b:raise HTTPException(404,"بسته واردسازی پیدا نشد.")
   if b["status"]!="PREVIEW":raise HTTPException(409,"این بسته قبلاً پردازش شده است.")
   rows=json.loads(b["rows_json"]);report=json.loads(b["report_json"]);ds=b["dataset"]
   if report["errors"]:raise HTTPException(422,"خطاهای پیش‌نمایش را اصلاح کنید.")
   n=0;opening=0
   try:
    for i,x in enumerate(rows,2):
     if ds=="customers":
      code=str(x["کد مشتری"]).strip();name=str(x["نام مشتری"]).strip();phone=str(x.get("موبایل") or x.get("تلفن") or "")
      cid=c.execute("INSERT INTO customers(tenant_id,name,phone) VALUES(?,?,?)",(t,name,phone)).lastrowid;c.execute("INSERT INTO dist_customer_codes(tenant_id,customer_id,code) VALUES(?,?,?)",(t,cid,code))
      bal=minor(x.get("مانده طلب افتتاحیه",0));opening+=bal
      if bal:c.execute("INSERT INTO dist_opening_balances(tenant_id,kind,code,name,amount_minor,batch_id) VALUES(?,?,?,?,?,?)",(t,"CUSTOMER",code,name,bal,bid))
     elif ds=="suppliers":
      code=str(x["کد تأمین‌کننده"]).strip();name=str(x["نام تأمین‌کننده"]).strip()
      c.execute("INSERT INTO dist_suppliers(tenant_id,code,name,phone,address) VALUES(?,?,?,?,?)",(t,code,name,str(x.get("موبایل") or x.get("تلفن") or ""),str(x.get("آدرس") or "")))
      bal=minor(x.get("مانده بدهی افتتاحیه",0));opening+=bal
      if bal:c.execute("INSERT INTO dist_opening_balances(tenant_id,kind,code,name,amount_minor,batch_id) VALUES(?,?,?,?,?,?)",(t,"SUPPLIER",code,name,bal,bid))
     elif ds=="warehouses":c.execute("INSERT INTO dist_warehouses(tenant_id,code,name,location) VALUES(?,?,?,?)",(t,str(x["کد انبار"]).strip(),str(x["نام انبار"]).strip(),str(x.get("نام محل") or x.get("محل بالادست") or "")))
     elif ds=="items":
      sku=str(x["کد کالا"]).strip();price=minor(x.get("قیمت فروش",0));cost=minor(x.get("قیمت خرید",0));qty=float(x.get("موجودی اولیه") or 0)
      if qty<0:raise ValueError(f"ردیف {i}: موجودی منفی")
      pid=c.execute("INSERT INTO products(tenant_id,sku,name,price,stock) VALUES(?,?,?,?,?)",(t,sku,str(x["نام کالا"]).strip(),price/100,qty)).lastrowid
      wh=str(x.get("انبار") or "DEFAULT");c.execute("INSERT INTO dist_item_data(tenant_id,product_id,cost_minor,min_stock,unit,warehouse_code) VALUES(?,?,?,?,?,?)",(t,pid,cost,float(x.get("حداقل موجودی") or 0),str(x.get("واحد") or "عدد"),wh))
      if qty:c.execute("INSERT INTO dist_stock_moves(tenant_id,product_id,warehouse_code,quantity,cost_minor,kind,reference) VALUES(?,?,?,?,?,?,?)",(t,pid,wh,qty,cost,"OPENING","XLSX"))
      opening+=minor(Decimal(str(qty))*Decimal(cost)/100)
     elif ds in ("banks","cash"):
      code=str(x.get("کد بانک") if ds=="banks" else x.get("کد صندوق")).strip();name=str(x.get("نام بانک") if ds=="banks" else x.get("نام صندوق")).strip();a=minor(x.get("مانده افتتاحیه",0));kind="BANK" if ds=="banks" else "CASH"
      c.execute("INSERT INTO dist_cash_accounts(tenant_id,code,name,kind,opening_minor) VALUES(?,?,?,?,?)",(t,code,name,kind,a));opening+=a
     elif ds in ("receivables","payables"):
      code=str(x["کد طرف حساب"]).strip();name=str(x["نام مشتری"] if ds=="receivables" else x["نام تأمین‌کننده"]).strip();a=minor(x.get("مبلغ طلب" if ds=="receivables" else "مبلغ بدهی",0));due=str(x.get("تاریخ سررسید") or "")
      if due:date.fromisoformat(due)
      kind="CUSTOMER" if ds=="receivables" else "SUPPLIER"
      if ds=="receivables" and not c.execute("SELECT 1 FROM dist_customer_codes WHERE tenant_id=? AND code=?",(t,code)).fetchone():raise ValueError(f"کد مشتری {code} ثبت نشده است.")
      if ds=="payables" and not c.execute("SELECT 1 FROM dist_suppliers WHERE tenant_id=? AND code=?",(t,code)).fetchone():raise ValueError(f"کد تأمین‌کننده {code} ثبت نشده است.")
      if c.execute("SELECT 1 FROM dist_opening_balances WHERE tenant_id=? AND kind=? AND code=?",(t,kind,code)).fetchone():raise ValueError(f"مانده افتتاحیه برای {code} قبلاً وارد شده است؛ از واردسازی تکراری خودداری کنید.")
      c.execute("INSERT INTO dist_opening_balances(tenant_id,kind,code,name,amount_minor,due_date,description,batch_id) VALUES(?,?,?,?,?,?,?,?)",(t,kind,code,name,a,due,str(x.get("شرح") or ""),bid));opening+=a
     elif ds=="opening_finance":
      code=str(x["کد حساب"]).strip();name=str(x["نام حساب"]).strip();d=minor(x.get("بدهکار افتتاحیه"));cr=minor(x.get("بستانکار افتتاحیه"))
      if d and cr:raise ValueError(f"ردیف {i}: بدهکار و بستانکار همزمان مجاز نیست.")
      if not c.execute("SELECT 1 FROM dist_accounts WHERE tenant_id=? AND code=?",(t,code)).fetchone():
       nature={"1":"asset","2":"liability","3":"equity","4":"income","5":"expense","6":"expense"}.get(code[:1],"expense")
       c.execute("INSERT INTO dist_accounts(tenant_id,code,name,nature) VALUES(?,?,?,?)",(t,code,name,nature))
     n+=1
    if ds in ("items","banks","cash","customers","suppliers","receivables","payables") and opening:
     if ds in ("customers","receivables"):lines=[{"code":"1100","d":opening},{"code":"3000","c":opening}]
     elif ds in ("suppliers","payables"):lines=[{"code":"3000","d":opening},{"code":"2000","c":opening}]
     else:lines=[{"code":"1200" if ds=="items" else "1000","d":opening},{"code":"3000","c":opening}]
     post(c,t,date.today().isoformat(),"مانده افتتاحیه واردشده از اکسل","XLSX_OPENING",None,lines,u.get("id"))
    if ds=="opening_finance":
     lines=[]
     for x in rows:
      d=minor(x.get("بدهکار افتتاحیه"));cr=minor(x.get("بستانکار افتتاحیه"));code=str(x["کد حساب"]).strip()
      if d:lines.append({"code":code,"d":d})
      if cr:lines.append({"code":code,"c":cr})
     if lines:post(c,t,date.today().isoformat(),"سند افتتاحیه مالی اکسل","XLSX_LEDGER",None,lines,u.get("id"))
    c.execute("UPDATE dist_import_batches SET status='COMMITTED',committed_at=CURRENT_TIMESTAMP WHERE id=?",(bid,));audit(c,t,f"distribution.import.committed:{ds}:{n}")
   except HTTPException:raise
   except Exception as e:raise HTTPException(409,f"واردسازی اتمیک انجام نشد؛ هیچ ردیفی ثبت نشد. ردیف/علت: {str(e)[:180]}") from e
  return {"status":"COMMITTED","batch_id":bid,"dataset":ds,"imported":n,"rejected":0}
 @r.get("/import/batches")
 def batches(u=Depends(user_dep)):
  with db() as c:rows=c.execute("SELECT id,dataset,status,report_json,created_at,committed_at FROM dist_import_batches WHERE tenant_id=? ORDER BY created_at DESC LIMIT 100",(tenant(u),)).fetchall()
  return [{"id":x["id"],"dataset":x["dataset"],"status":x["status"],"report":json.loads(x["report_json"]),"created_at":x["created_at"],"committed_at":x["committed_at"]} for x in rows]
 return r
